from __future__ import annotations

import zipfile

from docx import Document
from docx.oxml import parse_xml

from reviscope.ingest import ingest

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def _lines(document, tmp_path) -> list[str]:
    path = tmp_path / "paper.docx"
    document.save(path)
    return ingest(path).text.splitlines()


def test_caption_wrapped_table_keeps_caption_and_rows(tmp_path):
    # Quarto wraps each captioned table in a one-cell layout table.
    document = Document()
    document.add_paragraph("Intro text.")
    wrapper = document.add_table(rows=1, cols=1).cell(0, 0)
    wrapper.paragraphs[0].text = "Table 1: Synthetic estimates"
    data = wrapper.add_table(rows=2, cols=2)
    data.cell(0, 0).text, data.cell(0, 1).text = "Model", "Estimate"
    data.cell(1, 0).text, data.cell(1, 1).text = "Alpha model", "0.12"
    document.add_paragraph("After table.")
    assert _lines(document, tmp_path) == [
        "Intro text.", "Table 1: Synthetic estimates", "Model | Estimate", "Alpha model | 0.12", "After table."
    ]


def test_nested_tables_at_depth_are_emitted_in_order(tmp_path):
    document = Document()
    outer = document.add_table(rows=1, cols=2)
    outer.cell(0, 0).text = "left"
    cell = outer.cell(0, 1)
    cell.paragraphs[0].text = "before inner"
    inner = cell.add_table(rows=1, cols=1)
    inner.cell(0, 0).paragraphs[0].text = "inner"
    inner.cell(0, 0).add_table(rows=1, cols=2).rows[0].cells[1].text = "deepest"
    cell.add_paragraph("after inner")
    assert _lines(document, tmp_path) == ["left", "before inner", "inner", " | deepest", "after inner"]


def test_merged_cells_are_not_duplicated(tmp_path):
    document = Document()
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 2)).text = "Section heading"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "Tall"
    for row in (1, 2):
        table.cell(row, 1).text, table.cell(row, 2).text = f"b{row}", f"c{row}"
    lines = _lines(document, tmp_path)
    assert lines == ["Section heading", "Tall | b1 | c1", " | b2 | c2"]
    assert sum(line.count("Tall") for line in lines) == 1


def test_math_text_boxes_and_tracked_deletions(tmp_path):
    document = Document()
    paragraph = document.add_paragraph("Effect ")
    paragraph._p.append(parse_xml(
        '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
        "<m:r><m:t>Δr = 0.1</m:t></m:r></m:oMath>"
    ))
    paragraph._p.append(parse_xml(f"<w:del {W} w:id=\"1\" w:author=\"x\"><w:r><w:delText>gone</w:delText></w:r></w:del>"))
    paragraph._p.append(parse_xml(
        f'<w:r {W} xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        "<mc:AlternateContent><mc:Choice Requires=\"wps\"><w:drawing><wps:txbx><w:txbxContent>"
        "<w:p><w:r><w:t>Boxed note</w:t></w:r></w:p></w:txbxContent></wps:txbx></w:drawing></mc:Choice>"
        "<mc:Fallback><w:pict><w:txbxContent><w:p><w:r><w:t>Boxed note</w:t></w:r></w:p>"
        "</w:txbxContent></w:pict></mc:Fallback></mc:AlternateContent></w:r>"
    ))
    assert _lines(document, tmp_path) == ["Effect Δr = 0.1", "Boxed note"]


def test_footnotes_are_appended_with_reference_numbers(tmp_path):
    document = Document()
    paragraph = document.add_paragraph("Claim.")
    paragraph._p.append(parse_xml(f'<w:r {W}><w:footnoteReference w:id="7"/></w:r>'))
    path = tmp_path / "plain.docx"
    document.save(path)
    footnotes = (
        f'<w:footnotes {W}><w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
        '<w:footnote w:id="7"><w:p><w:r><w:footnoteRef/></w:r><w:r><w:t>Synthetic note text.</w:t></w:r></w:p></w:footnote>'
        "</w:footnotes>"
    )
    noted = tmp_path / "noted.docx"
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(noted, "w") as dst:
        for item in src.infolist():
            data = src.read(item)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"</Types>", b'<Override PartName="/word/footnotes.xml" ContentType="application/'
                                    b'vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>')
            elif item.filename == "word/_rels/document.xml.rels":
                data = data.replace(b"</Relationships>", b'<Relationship Id="rIdFn" Type="http://schemas.openxmlformats.org/'
                                    b'officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/></Relationships>')
            dst.writestr(item, data)
        dst.writestr("word/footnotes.xml", footnotes)
    assert ingest(noted).text.splitlines() == ["Claim.[footnote 1]", "", "Footnotes", "[Footnote 1] Synthetic note text."]
