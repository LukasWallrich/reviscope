from __future__ import annotations

from pathlib import Path

from docx import Document
import pytest
from pydantic import BaseModel

from reviscope.backend import Backend, FixtureBackend
from reviscope.ingest import ingest
from reviscope.pipeline import ReviewPipeline
from reviscope.render import to_html
from reviscope.render import to_markdown
from reviscope.schemas import Evidence, Finding, Profile, StudyMap


class ScriptedBackend(Backend):
    """Typed pipeline fake: behavior is explicit for every response schema."""

    name = "scripted"

    def __init__(self, *, model="m1", effort="low", modules=None, verify=None, editorial=None):
        self.model, self.effort = model, effort
        self.modules = modules or {}
        self.verify = verify
        self.editorial = editorial
        self.calls = []

    def generate(self, instruction, evidence, response_model):
        self.calls.append((response_model.__name__, instruction))
        if response_model.__name__ == "StudyMap":
            return StudyMap(studies=[], research_question="q", design_summary="d", contribution_summary="c", strengths=[])
        if response_model.__name__ == "FindingsResponse":
            module = next((key for key in self.modules if key in instruction), "default")
            return response_model.model_validate({"findings": self.modules.get(module, [])})
        if response_model.__name__ == "VerificationResponse":
            rows = self.verify(instruction) if callable(self.verify) else (self.verify or [])
            return response_model.model_validate({"decisions": rows})
        if response_model.__name__ == "EditorialResponse":
            rows = self.editorial(instruction) if callable(self.editorial) else (self.editorial or [])
            return response_model.model_validate({"decisions": rows, "reconciled_overview": {
                "design_summary": "d", "contribution_summary": "c", "strengths": []}})
        raise AssertionError(response_model)


def profile(*modules):
    return Profile(id="test", title="Test", modules=list(modules),
                   module_prompts={item: f"RUN {item}" for item in modules},
                   verification_prompt="verify", editorial_prompt="edit")


def manuscript(tmp_path, text="Quoted source sentence. " + "Context for a complete manuscript. " * 40):
    path = tmp_path / "paper.md"
    path.write_text(text)
    return path


def finding(identifier="same", claim="A grounded concern", status="recomputed"):
    return Finding(id=identifier, module="forged", claim=claim, rationale="r", remedy="fix",
                   status=status, evidence=[Evidence(source_id="placeholder", quote="Quoted source sentence.")]).model_dump()


def support_all(instruction):
    import json, re
    block = instruction.split("CANDIDATES\n", 1)[1].split("\nDISCIPLINE RULES", 1)[0]
    candidates = json.loads(block)
    return [{"finding_id": item["finding_id"], "status": "supported", "rationale": "checked",
             "evidence": [{"source_id": item["quoted_evidence"][0]["source_id"], "quote": "Quoted source sentence."}]}
            for item in candidates]


def _fix_source_ids(rows, source_id):
    for row in rows:
        row["evidence"][0]["source_id"] = source_id
    return rows


def test_model_cannot_forge_recomputed_status_and_duplicate_ids_are_namespaced(tmp_path):
    paper = manuscript(tmp_path)
    source_id = ingest(paper).id
    rows_a = _fix_source_ids([finding()], source_id)
    rows_b = _fix_source_ids([finding()], source_id)
    generator = ScriptedBackend(modules={"RUN a": rows_a, "RUN b": rows_b}, editorial=[])
    verifier = ScriptedBackend(model="verifier", verify=support_all)
    run = ReviewPipeline(generator, profile("a", "b"), verifier).run(paper, output_dir=tmp_path / "out")
    assert len({item.id for item in run.candidates}) == 2
    assert all(item.status == "candidate" for item in run.candidates)
    assert all(item.status == "llm_supported" for item in run.findings)


def test_corrected_verifier_quotes_reach_editorial_and_published_output(tmp_path):
    import json

    paper = manuscript(tmp_path)
    rows = _fix_source_ids([finding()], ingest(paper).id)
    rows[0]["evidence"][0]["quote"] = "Invented generator quotation."

    def editorial(instruction):
        findings = json.loads(instruction.split("\nFINDINGS\n", 1)[1])
        assert findings[0]["status"] == "llm_supported"
        assert findings[0]["evidence"][0]["quote"] == "Quoted source sentence."
        return [{"finding_id": findings[0]["id"], "disposition": "keep", "reason": "Supported"}]

    generator = ScriptedBackend(modules={"RUN a": rows}, editorial=editorial)
    run = ReviewPipeline(generator, profile("a"), ScriptedBackend(verify=support_all)).run(
        paper, output_dir=tmp_path / "out")
    assert not run.partial
    assert run.findings[0].editorial_disposition == "publish"
    assert run.candidates[0].evidence[0].quote == "Invented generator quotation."
    published = to_markdown(run).split("## Coverage and audit", 1)[0]
    assert "Quoted source sentence." in published
    assert "Invented generator quotation." not in published


@pytest.mark.parametrize("mode", ["unknown", "duplicate", "missing"])
def test_bad_verifier_join_fails_closed_and_marks_partial(tmp_path, mode):
    paper = manuscript(tmp_path)
    source_id = ingest(paper).id
    generator = ScriptedBackend(modules={"RUN a": _fix_source_ids([finding()], source_id)})
    valid = {"finding_id": "a:0:same", "status": "supported", "rationale": "x",
             "evidence": [{"source_id": source_id, "quote": "Quoted source sentence."}]}
    bad = ([] if mode == "missing" else
           [{**valid, "finding_id": "unknown"}] if mode == "unknown" else
           [valid, valid])
    run = ReviewPipeline(generator, profile("a"), ScriptedBackend(verify=bad)).run(
        paper, output_dir=tmp_path / "out")
    assert run.partial
    assert next(stage for stage in run.stages if stage.name == "verification").status != "completed"
    assert all(item.status in {"unverified", "unresolved"} for item in run.findings)


def test_editorial_cannot_eliminate_only_supported_duplicate_via_rejected_target(tmp_path):
    paper = manuscript(tmp_path)
    source_id = ingest(paper).id
    rows = _fix_source_ids([finding()], source_id)
    generator = ScriptedBackend(modules={"RUN a": rows, "RUN b": rows}, verify=support_all,
        editorial=lambda instruction: [
            {"finding_id": "a:0:same", "disposition": "reject", "reason": "reject", "target_id": None},
            {"finding_id": "b:0:same", "disposition": "merge", "reason": "duplicate", "target_id": "a:0:same"},
        ])
    run = ReviewPipeline(generator, profile("a", "b"), generator).run(paper, output_dir=tmp_path / "out")
    assert run.partial or any(item.editorial_disposition == "publish" for item in run.findings)


def test_overreaching_remedy_is_withheld_without_discarding_supported_claim(tmp_path):
    paper = manuscript(tmp_path)
    source_id = ingest(paper).id
    rows = _fix_source_ids([finding()], source_id)
    def verifier(instruction):
        decisions = support_all(instruction)
        for decision in decisions:
            decision.update({"remedy_status": "overreaching", "remedy_rationale": "The requested fix exceeds the evidenced concern."})
        return decisions
    generator = ScriptedBackend(modules={"RUN a": rows}, editorial=lambda _instruction: [
        {"finding_id": "a:0:same", "disposition": "keep", "reason": "Valid claim", "target_id": None}
    ])
    run = ReviewPipeline(generator, profile("a"), ScriptedBackend(model="verifier", verify=verifier)).run(
        paper, output_dir=tmp_path / "out")
    item = run.findings[0]
    assert item.status == "llm_supported" and item.editorial_disposition == "publish"
    assert item.remedy == "fix"
    rendered = to_markdown(run)
    assert "Proposed response withheld" in rendered
    assert "**Suggested response:** fix" not in rendered


def test_cache_identity_changes_with_backend_model_effort_profile_and_input(tmp_path):
    out = tmp_path / "out"
    paper = manuscript(tmp_path)
    first = ScriptedBackend(model="m1", effort="low")
    ReviewPipeline(first, profile()).run(paper, output_dir=out)
    again = ScriptedBackend(model="m1", effort="low")
    ReviewPipeline(again, profile()).run(paper, output_dir=out)
    assert not any(name == "StudyMap" for name, _ in again.calls)
    changed = ScriptedBackend(model="m2", effort="high")
    ReviewPipeline(changed, Profile(id="other", title="Other", modules=[], metadata={"revision": 2})).run(
        paper, output_dir=out)
    assert any(name == "StudyMap" for name, _ in changed.calls)
    paper.write_text("Changed input.")
    changed_input = ScriptedBackend(model="m2", effort="high")
    ReviewPipeline(changed_input, profile()).run(paper, output_dir=out)
    assert any(name == "StudyMap" for name, _ in changed_input.calls)

    class SchemaA(BaseModel):
        value: int
    class SchemaB(BaseModel):
        value: int
        note: str = ""
    cache_pipeline = ReviewPipeline(ScriptedBackend(), profile())
    calls = []
    cache_pipeline._cached(out, "contract", {"upstream": "a"}, SchemaA,
                            lambda: calls.append("first") or SchemaA(value=1))
    cache_pipeline._cached(out, "contract", {"upstream": "a"}, SchemaA,
                            lambda: calls.append("cached") or SchemaA(value=1))
    cache_pipeline._cached(out, "contract", {"upstream": "b"}, SchemaA,
                            lambda: calls.append("upstream") or SchemaA(value=1))
    cache_pipeline._cached(out, "contract", {"upstream": "b"}, SchemaB,
                            lambda: calls.append("schema") or SchemaB(value=1))
    assert calls == ["first", "upstream", "schema"]


def test_html_escapes_manuscript_script_and_marks_fixture_demo(tmp_path):
    paper = manuscript(tmp_path, "<script>alert('x')</script> Participants were recruited from the university pool.")
    ReviewPipeline(FixtureBackend()).run(paper, output_dir=tmp_path / "out")
    rendered = (tmp_path / "out" / "review.html").read_text()
    escaped = to_html("Manuscript says <script>alert('x')</script>")
    assert "<script>alert" not in escaped
    assert "&lt;script&gt;" in escaped
    assert "fixture" in rendered and "deterministic-demo" in rendered


def test_pdf_blank_page_warning_and_docx_table_order(monkeypatch, tmp_path):
    class Page:
        def __init__(self, text): self.text = text
        def extract_text(self): return self.text
    class Reader:
        def __init__(self, _): self.pages = [Page("first"), Page("")]
    import pypdf
    monkeypatch.setattr(pypdf, "PdfReader", Reader)
    pdf = tmp_path / "paper.pdf"
    pdf.write_bytes(b"fake")
    extracted = ingest(pdf)
    assert extracted.pages[1].text == ""
    assert "page(s): 2" in extracted.extraction_warnings[0]
    assert ReviewPipeline(FixtureBackend()).run(pdf, output_dir=tmp_path / "pdf-out").partial

    document = Document()
    document.add_paragraph("before")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "left", "right"
    document.add_paragraph("after")
    docx = tmp_path / "paper.docx"
    document.save(docx)
    assert ingest(docx).text.splitlines() == ["before", "left | right", "after"]
